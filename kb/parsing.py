"""Document parsing for PDF, DOCX, TXT, MD.

Returns ParsedDocument with per-page text so we can preserve page numbers as
citation metadata. PDFs preserve real page numbers; DOCX/TXT/MD have a single
synthetic page (None) since they are unpaginated at the source level.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from pathlib import Path

from common.citations import DocumentType


@dataclass
class ParsedPage:
    page: int | None  # 1-indexed for PDFs; None for unpaginated formats
    text: str


@dataclass
class ParsedDocument:
    filename: str
    document_type: DocumentType
    sha256: str
    pages: list[ParsedPage] = field(default_factory=list)


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _read_pdf(path: Path) -> list[ParsedPage]:
    from pypdf import PdfReader  # lazy import; not needed for plain-text tests

    reader = PdfReader(str(path))
    out: list[ParsedPage] = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            out.append(ParsedPage(page=i, text=text))
    return out


def _read_docx(path: Path) -> list[ParsedPage]:
    from docx import Document as DocxDocument  # lazy import

    doc = DocxDocument(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    text = "\n".join(paragraphs).strip()
    return [ParsedPage(page=None, text=text)] if text else []


def _read_text(path: Path) -> list[ParsedPage]:
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    return [ParsedPage(page=None, text=text)] if text else []


def parse_document(path: str | Path, document_type: DocumentType = "other") -> ParsedDocument:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(p)

    raw = p.read_bytes()
    digest = _sha256(raw)

    suffix = p.suffix.lower()
    if suffix == ".pdf":
        pages = _read_pdf(p)
    elif suffix == ".docx":
        pages = _read_docx(p)
    elif suffix in {".txt", ".md"}:
        pages = _read_text(p)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")

    return ParsedDocument(
        filename=p.name,
        document_type=document_type,
        sha256=digest,
        pages=pages,
    )
